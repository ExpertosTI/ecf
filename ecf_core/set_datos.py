"""Convierte el Excel de set de datos del portal DGII en XML e-CF y RFCE.

El portal marca las celdas vacías con ``#e``. Cada fila no vacía se arma en el
orden del XSD del tipo y se valida en local. No envía nada a la DGII.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import datetime
from decimal import Decimal
from pathlib import Path

from lxml import etree
from openpyxl import load_workbook

from ecf_core.ecf_core_service import XSD_DIR

_VACIO = {"", "#e", "#E"}
_INDEX = re.compile(r"^(.+?)\[(\d+)\](?:\[(\d+)\])?$")

_INFO_EXPORTACION = {
    "NombrePuertoEmbarque", "CondicionesEntrega", "TotalFob", "Seguro", "Flete",
    "TotalCif", "RegimenAduanero", "NombrePuertoSalida", "NombrePuertoDesembarque",
}
_TRANSPORTE_EXPORTACION = (
    "ViaTransporte", "PaisOrigen", "DireccionDestino", "PaisDestino",
)
_COMPRADOR_EXTRANJERO = {33, 34, 41, 43, 46, 47}


@dataclass
class FilaSet:
    caso: str
    tipo: int
    encf: str
    destino: str
    xml: bytes
    errores: list[str] = field(default_factory=list)


def _texto(valor) -> str | None:
    if valor is None:
        return None
    if isinstance(valor, datetime):
        return valor.strftime("%d-%m-%Y")
    if isinstance(valor, float):
        if valor.is_integer():
            return str(int(valor))
        return f"{valor:.2f}"
    texto = str(valor).strip()
    if texto in _VACIO:
        return None
    return texto


def _dinero(valor) -> str | None:
    texto = _texto(valor)
    if texto is None:
        return None
    try:
        return f"{Decimal(texto):.2f}"
    except Exception:
        return texto


class _Fila:
    def __init__(self, header: list[str], values: tuple):
        self._data: dict[str, object] = {}
        for nombre, valor in zip(header, values):
            if not nombre:
                continue
            self._data[str(nombre).strip()] = valor

    def get(self, nombre: str, dinero: bool = False) -> str | None:
        if nombre not in self._data:
            return None
        return _dinero(self._data[nombre]) if dinero else _texto(self._data[nombre])

    def linea(self, nombre: str, n: int, sub: int | None = None, dinero: bool = False) -> str | None:
        clave = f"{nombre}[{n}]" if sub is None else f"{nombre}[{n}][{sub}]"
        # El Excel escribe «MontosubRecargo» en lugar de «MontoSubRecargo».
        if self.get(clave, dinero=dinero) is None and nombre == "MontoSubRecargo":
            return self.get(f"MontosubRecargo[{n}][{sub}]", dinero=dinero)
        return self.get(clave, dinero=dinero)


def _e(parent, tag: str, text: str | None):
    if text is None:
        return None
    nodo = etree.SubElement(parent, tag)
    nodo.text = text
    return nodo


def _tiene(nodo) -> bool:
    return len(nodo) > 0


def _formas_pago(parent, fila: _Fila):
    tabla = etree.Element("TablaFormasPago")
    for i in range(1, 8):
        forma = fila.get(f"FormaPago[{i}]")
        monto = fila.get(f"MontoPago[{i}]", dinero=True)
        if forma is None and monto is None:
            continue
        item = etree.SubElement(tabla, "FormaDePago")
        _e(item, "FormaPago", forma)
        _e(item, "MontoPago", monto)
    if _tiene(tabla):
        parent.append(tabla)


def _id_doc(enc, fila: _Fila, tipo: int):
    id_doc = etree.SubElement(enc, "IdDoc")
    _e(id_doc, "TipoeCF", fila.get("TipoeCF"))
    _e(id_doc, "eNCF", fila.get("ENCF"))
    if tipo == 34:
        _e(id_doc, "IndicadorNotaCredito", fila.get("IndicadorNotaCredito") or "1")
    elif tipo != 32:
        _e(id_doc, "FechaVencimientoSecuencia", fila.get("FechaVencimientoSecuencia"))
    _e(id_doc, "IndicadorEnvioDiferido", fila.get("IndicadorEnvioDiferido"))
    _e(id_doc, "IndicadorMontoGravado", fila.get("IndicadorMontoGravado"))
    _e(id_doc, "TipoIngresos", fila.get("TipoIngresos"))
    _e(id_doc, "TipoPago", fila.get("TipoPago"))
    _e(id_doc, "FechaLimitePago", fila.get("FechaLimitePago"))
    _e(id_doc, "TerminoPago", fila.get("TerminoPago"))
    _formas_pago(id_doc, fila)
    _e(id_doc, "TipoCuentaPago", fila.get("TipoCuentaPago"))
    _e(id_doc, "NumeroCuentaPago", fila.get("NumeroCuentaPago"))
    _e(id_doc, "BancoPago", fila.get("BancoPago"))
    _e(id_doc, "FechaDesde", fila.get("FechaDesde"))
    _e(id_doc, "FechaHasta", fila.get("FechaHasta"))
    _e(id_doc, "TotalPaginas", fila.get("TotalPaginas"))


def _emisor(enc, fila: _Fila):
    emisor = etree.SubElement(enc, "Emisor")
    for tag in (
        "RNCEmisor", "RazonSocialEmisor", "NombreComercial", "Sucursal",
        "DireccionEmisor", "Municipio", "Provincia",
    ):
        _e(emisor, tag, fila.get(tag))
    telefonos = etree.Element("TablaTelefonoEmisor")
    for i in range(1, 4):
        _e(telefonos, "TelefonoEmisor", fila.get(f"TelefonoEmisor[{i}]"))
    if _tiene(telefonos):
        emisor.append(telefonos)
    for tag in (
        "CorreoEmisor", "WebSite", "ActividadEconomica", "CodigoVendedor",
        "NumeroFacturaInterna", "NumeroPedidoInterno", "ZonaVenta", "RutaVenta",
        "InformacionAdicionalEmisor", "FechaEmision",
    ):
        _e(emisor, tag, fila.get(tag))


def _comprador(enc, fila: _Fila, tipo: int):
    comprador = etree.Element("Comprador")
    _e(comprador, "RNCComprador", fila.get("RNCComprador"))
    if tipo in _COMPRADOR_EXTRANJERO:
        _e(comprador, "IdentificadorExtranjero", fila.get("IdentificadorExtranjero"))
    for tag in (
        "RazonSocialComprador", "ContactoComprador", "CorreoComprador",
        "DireccionComprador", "MunicipioComprador", "ProvinciaComprador",
        "PaisComprador", "FechaEntrega", "ContactoEntrega", "DireccionEntrega",
        "TelefonoAdicional", "FechaOrdenCompra", "NumeroOrdenCompra",
        "CodigoInternoComprador", "ResponsablePago", "InformacionAdicionalComprador",
    ):
        _e(comprador, tag, fila.get(tag))
    if _tiene(comprador):
        enc.append(comprador)


def _informaciones(enc, fila: _Fila, tipo: int):
    info = etree.Element("InformacionesAdicionales")
    campos = [
        "FechaEmbarque", "NumeroEmbarque", "NumeroContenedor", "NumeroReferencia",
    ]
    if tipo == 46:
        campos += [
            "NombrePuertoEmbarque", "CondicionesEntrega", "TotalFob", "Seguro",
            "Flete", "TotalCif", "RegimenAduanero", "NombrePuertoSalida",
            "NombrePuertoDesembarque",
        ]
    campos += [
        "PesoBruto", "PesoNeto", "UnidadPesoBruto", "UnidadPesoNeto",
        "CantidadBulto", "UnidadBulto", "VolumenBulto", "UnidadVolumen",
    ]
    for tag in campos:
        if tag in _INFO_EXPORTACION and tipo != 46:
            continue
        dinero = tag in {"TotalFob", "Seguro", "Flete", "TotalCif", "PesoBruto", "PesoNeto", "CantidadBulto", "VolumenBulto"}
        _e(info, tag, fila.get(tag, dinero=dinero))
    if _tiene(info):
        enc.append(info)


def _transporte(enc, fila: _Fila, tipo: int):
    transporte = etree.Element("Transporte")
    if tipo == 46:
        for tag in _TRANSPORTE_EXPORTACION:
            _e(transporte, tag, fila.get(tag))
    for tag in (
        "Conductor", "DocumentoTransporte", "Ficha", "Placa",
        "RutaTransporte", "ZonaTransporte", "NumeroAlbaran",
    ):
        _e(transporte, tag, fila.get(tag))
    if _tiene(transporte):
        enc.append(transporte)


def _totales(enc, fila: _Fila):
    totales = etree.SubElement(enc, "Totales")
    for tag in (
        "MontoGravadoTotal", "MontoGravadoI1", "MontoGravadoI2", "MontoGravadoI3",
        "MontoExento",
    ):
        _e(totales, tag, fila.get(tag, dinero=True))
    for tag in ("ITBIS1", "ITBIS2", "ITBIS3"):
        _e(totales, tag, fila.get(tag))
    for tag in (
        "TotalITBIS", "TotalITBIS1", "TotalITBIS2", "TotalITBIS3",
        "MontoImpuestoAdicional", "MontoTotal", "MontoNoFacturable",
        "MontoPeriodo", "SaldoAnterior", "MontoAvancePago", "ValorPagar",
        "TotalITBISRetenido", "TotalISRRetencion",
    ):
        _e(totales, tag, fila.get(tag, dinero=True))


def _item(detalles, fila: _Fila, n: int):
    if fila.linea("NumeroLinea", n) is None and fila.linea("NombreItem", n) is None:
        return
    item = etree.SubElement(detalles, "Item")
    _e(item, "NumeroLinea", fila.linea("NumeroLinea", n))
    codigos = etree.Element("TablaCodigosItem")
    for sub in range(1, 6):
        tipo_cod = fila.linea("TipoCodigo", n, sub)
        codigo = fila.linea("CodigoItem", n, sub)
        if tipo_cod is None and codigo is None:
            continue
        cod = etree.SubElement(codigos, "CodigosItem")
        _e(cod, "TipoCodigo", tipo_cod)
        _e(cod, "CodigoItem", codigo)
    if _tiene(codigos):
        item.append(codigos)
    _e(item, "IndicadorFacturacion", fila.linea("IndicadorFacturacion", n))
    retencion = etree.Element("Retencion")
    _e(retencion, "IndicadorAgenteRetencionoPercepcion", fila.linea("IndicadorAgenteRetencionoPercepcion", n))
    _e(retencion, "MontoITBISRetenido", fila.linea("MontoITBISRetenido", n, dinero=True))
    _e(retencion, "MontoISRRetenido", fila.linea("MontoISRRetenido", n, dinero=True))
    if _tiene(retencion):
        item.append(retencion)
    _e(item, "NombreItem", fila.linea("NombreItem", n))
    _e(item, "IndicadorBienoServicio", fila.linea("IndicadorBienoServicio", n))
    _e(item, "DescripcionItem", fila.linea("DescripcionItem", n))
    _e(item, "CantidadItem", fila.linea("CantidadItem", n, dinero=True))
    _e(item, "UnidadMedida", fila.linea("UnidadMedida", n))
    _e(item, "FechaElaboracion", fila.linea("FechaElaboracion", n))
    _e(item, "FechaVencimientoItem", fila.linea("FechaVencimientoItem", n))
    _e(item, "PrecioUnitarioItem", fila.linea("PrecioUnitarioItem", n, dinero=True))
    _e(item, "DescuentoMonto", fila.linea("DescuentoMonto", n, dinero=True))
    descuentos = etree.Element("TablaSubDescuento")
    for sub in range(1, 13):
        tipo_d = fila.linea("TipoSubDescuento", n, sub)
        porc = fila.linea("SubDescuentoPorcentaje", n, sub, dinero=True)
        monto = fila.linea("MontoSubDescuento", n, sub, dinero=True)
        if tipo_d is None and porc is None and monto is None:
            continue
        sub_n = etree.SubElement(descuentos, "SubDescuento")
        _e(sub_n, "TipoSubDescuento", tipo_d)
        _e(sub_n, "SubDescuentoPorcentaje", porc)
        _e(sub_n, "MontoSubDescuento", monto)
    if _tiene(descuentos):
        item.append(descuentos)
    _e(item, "RecargoMonto", fila.linea("RecargoMonto", n, dinero=True))
    recargos = etree.Element("TablaSubRecargo")
    for sub in range(1, 13):
        tipo_r = fila.linea("TipoSubRecargo", n, sub)
        porc = fila.linea("SubRecargoPorcentaje", n, sub, dinero=True)
        monto = fila.linea("MontoSubRecargo", n, sub, dinero=True)
        if tipo_r is None and porc is None and monto is None:
            continue
        sub_n = etree.SubElement(recargos, "SubRecargo")
        _e(sub_n, "TipoSubRecargo", tipo_r)
        _e(sub_n, "SubRecargoPorcentaje", porc)
        _e(sub_n, "MontoSubRecargo", monto)
    if _tiene(recargos):
        item.append(recargos)
    otra = etree.Element("OtraMonedaDetalle")
    _e(otra, "PrecioOtraMoneda", fila.linea("PrecioOtraMoneda", n, dinero=True))
    _e(otra, "MontoItemOtraMoneda", fila.linea("MontoItemOtraMoneda", n, dinero=True))
    if _tiene(otra):
        item.append(otra)
    _e(item, "MontoItem", fila.linea("MontoItem", n, dinero=True))


def construir_ecf(fila: _Fila) -> bytes:
    tipo = int(fila.get("TipoeCF") or "0")
    root = etree.Element("ECF")
    enc = etree.SubElement(root, "Encabezado")
    _e(enc, "Version", fila.get("Version") or "1.0")
    _id_doc(enc, fila, tipo)
    _emisor(enc, fila)
    _comprador(enc, fila, tipo)
    _informaciones(enc, fila, tipo)
    _transporte(enc, fila, tipo)
    _totales(enc, fila)
    detalles = etree.SubElement(root, "DetallesItems")
    for n in range(1, 65):
        _item(detalles, fila, n)
    referencia = etree.Element("InformacionReferencia")
    _e(referencia, "NCFModificado", fila.get("NCFModificado"))
    _e(referencia, "FechaNCFModificado", fila.get("FechaNCFModificado"))
    _e(referencia, "CodigoModificacion", fila.get("CodigoModificacion"))
    _e(referencia, "RazonModificacion", fila.get("RazonModificacion"))
    if _tiene(referencia):
        root.append(referencia)
    _e(root, "FechaHoraFirma", datetime.now().strftime("%d-%m-%Y %H:%M:%S"))
    etree.SubElement(root, "{http://www.w3.org/2000/09/xmldsig#}Signature")
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8")


def construir_rfce(fila: _Fila) -> bytes:
    root = etree.Element("RFCE")
    enc = etree.SubElement(root, "Encabezado")
    _e(enc, "Version", fila.get("Version") or "1.0")
    id_doc = etree.SubElement(enc, "IdDoc")
    _e(id_doc, "TipoeCF", fila.get("TipoeCF"))
    _e(id_doc, "eNCF", fila.get("ENCF"))
    _e(id_doc, "TipoIngresos", fila.get("TipoIngresos"))
    _e(id_doc, "TipoPago", fila.get("TipoPago"))
    _formas_pago(id_doc, fila)
    emisor = etree.SubElement(enc, "Emisor")
    _e(emisor, "RNCEmisor", fila.get("RNCEmisor"))
    _e(emisor, "RazonSocialEmisor", fila.get("RazonSocialEmisor"))
    _e(emisor, "FechaEmision", fila.get("FechaEmision"))
    comprador = etree.Element("Comprador")
    _e(comprador, "RNCComprador", fila.get("RNCComprador"))
    _e(comprador, "IdentificadorExtranjero", fila.get("IdentificadorExtranjero"))
    _e(comprador, "RazonSocialComprador", fila.get("RazonSocialComprador"))
    if _tiene(comprador):
        enc.append(comprador)
    totales = etree.SubElement(enc, "Totales")
    for tag in (
        "MontoGravadoTotal", "MontoGravadoI1", "MontoGravadoI2", "MontoGravadoI3",
        "MontoExento", "TotalITBIS", "TotalITBIS1", "TotalITBIS2", "TotalITBIS3",
        "MontoImpuestoAdicional", "MontoTotal",
    ):
        _e(totales, tag, fila.get(tag, dinero=True))
    # Sale del e-CF firmado. En la validación local, si el Excel no lo trae,
    # se usa un marcador de 6 caracteres para poder cerrar el XSD.
    _e(enc, "CodigoSeguridadeCF", fila.get("CodigoSeguridadeCF") or "PEND01")
    etree.SubElement(root, "{http://www.w3.org/2000/09/xmldsig#}Signature")
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8")


def _validar(xml: bytes, nombre_xsd: str) -> list[str]:
    ruta = XSD_DIR / f"{nombre_xsd}.xsd"
    if not ruta.exists():
        return [f"XSD no encontrado: {ruta.name}"]
    schema = etree.XMLSchema(etree.parse(str(ruta)))
    doc = etree.fromstring(xml)
    if schema.validate(doc):
        return []
    return [err.message for err in schema.error_log]


def _destino(tipo: int, monto: str | None) -> str:
    if tipo == 32 and monto is not None:
        try:
            if Decimal(monto) < Decimal("250000"):
                return "consumo_menor"
        except Exception:
            pass
    return "recepcion"


def _filas_hoja(wb, nombre: str) -> list[_Fila]:
    if nombre not in wb.sheetnames:
        return []
    ws = wb[nombre]
    rows = ws.iter_rows(values_only=True)
    header = [str(c).strip() if c else "" for c in next(rows)]
    salida = []
    for values in rows:
        if not values or _texto(values[0]) is None:
            continue
        salida.append(_Fila(header, values))
    return salida


def cargar_set(path: str | Path) -> tuple[list[FilaSet], list[FilaSet]]:
    """Lee el Excel del portal y valida cada XML contra el XSD del tipo."""
    wb = load_workbook(path, read_only=True, data_only=True)
    ecf_rows = []
    for fila in _filas_hoja(wb, "ECF"):
        tipo = int(fila.get("TipoeCF") or "0")
        xml = construir_ecf(fila)
        ecf_rows.append(FilaSet(
            caso=fila.get("CasoPrueba") or "",
            tipo=tipo,
            encf=fila.get("ENCF") or "",
            destino=_destino(tipo, fila.get("MontoTotal", dinero=True)),
            xml=xml,
            errores=_validar(xml, f"ECF-{tipo}"),
        ))
    rfce_rows = []
    for fila in _filas_hoja(wb, "RFCE"):
        tipo = int(fila.get("TipoeCF") or "32")
        xml = construir_rfce(fila)
        rfce_rows.append(FilaSet(
            caso=fila.get("CasoPrueba") or "",
            tipo=tipo,
            encf=fila.get("ENCF") or "",
            destino="resumen",
            xml=xml,
            errores=_validar(xml, "RFCE-32"),
        ))
    wb.close()
    return ecf_rows, rfce_rows
